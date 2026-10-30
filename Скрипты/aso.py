"""Акты о смонтированном оборудовании (по томам) с ведомостью — python-docx."""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import common as C
from equip import EQ

import paths
OUT = str(paths.TMP)
NORM = {'СПС': 'СП 76.13330.2016 «Электротехнические устройства»; СП 6.13130.2025; СП 484.1311500.2020; ГОСТ Р 59638-2021',
        'СКУД': 'СП 76.13330.2016 «Электротехнические устройства»',
        'СКС': 'СП 76.13330.2016 «Электротехнические устройства»; ГОСТ Р 53245-2008',
        'СОТ': 'СП 76.13330.2016 «Электротехнические устройства»; ГОСТ Р 51558-2014',
        'СОТС': 'СП 76.13330.2016 «Электротехнические устройства»'}

SYSG = {'СПС': 'системы пожарной сигнализации и системы оповещения и управления эвакуацией (АПС, СОУЭ)',
        'СКУД': 'системы контроля и управления доступом (СКУД)', 'СКС': 'структурированной кабельной системы (СКС)',
        'СОТ': 'системы охранной телевизионной (СОТ)', 'СОТС': 'системы охранной и тревожной сигнализации (СОТС)'}

def font(run, sz=11, b=False, i=False):
    run.font.name = 'Times New Roman'; run.font.size = Pt(sz); run.bold = b; run.italic = i
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

def para(doc, text='', sz=11, b=False, i=False, align=None, after=2, before=0):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if align: p.alignment = align
    if text:
        font(p.add_run(text), sz, b, i)
    return p

def lab(doc, label, text, sz=10):
    p = para(doc)
    font(p.add_run(label + ' '), sz, True); font(p.add_run(text), sz, False, True)
    return p

def borders(table):
    tbl = table._tbl; tblPr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), '000000'); b.append(x)
    tblPr.append(b)

def fix_widths(table, widths):
    tblPr = table._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(int(w.twips if hasattr(w,'twips') else w)))
    for row in table.rows:
        for c, w in zip(row.cells, widths): c.width = w

def cell(c, text, sz=9, b=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    c.text = ''; p = c.paragraphs[0]; p.alignment = align; p.paragraph_format.space_after = Pt(0)
    font(p.add_run(str(text)), sz, b)

def build(bld, code):
    v = C.vol(bld, code); d = C.DATES[bld]
    doc = Document()
    s = doc.sections[0]; s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = Mm(25); s.right_margin = Mm(10); s.top_margin = Mm(15); s.bottom_margin = Mm(15)
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(10)
    lab(doc, 'Объект капитального строительства:', C.OBJ)
    lab(doc, 'Застройщик (технический заказчик):', C.ZAK)
    lab(doc, 'Лицо, осуществляющее строительство:', C.STR)
    lab(doc, 'Лицо, выполнившее монтаж:', C.ISP)
    no = f'{code}-{bld}-О'
    para(doc, 'АКТ', 14, True, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=0)
    para(doc, 'о смонтированном оборудовании', 12, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(doc, f'{C.SYSN[code]}, здание {bld}', 12, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell(t.rows[0].cells[0], f'№ {no}', 11); cell(t.rows[0].cells[1], f'«{d["act"][:2]}» {["", "января","февраля","марта","апреля","мая","июня","июля","августа","сентября","октября","ноября","декабря"][int(d["act"][3:5])]} {d["act"][6:]} г.', 11, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, 'Комиссия в составе:', 10, True, before=4)
    for role, who in [('Представитель застройщика (технического заказчика) по вопросам строительного контроля',
                       'главный специалист отдела строительного контроля Департамента строительного консалтинга Коренев Д.Ю., НРС НОСТРОЙ С-50-002782, приказ № 37 от 27.04.2026 г.'),
                      ('Представитель лица, осуществляющего строительство',
                       'руководитель проекта Белоногов И.Ю., НРС НОСТРОЙ С-72-258451, приказ № 2 от 14.01.2026 г.'),
                      ('Представитель лица, выполнившего монтаж (лица, осуществляющего подготовку проектной документации)',
                       'начальник отдела систем технической безопасности ООО «АльтЭн» Иванчук М.В., приказ № 346 от 01.09.2026 г.')]:
        lab(doc, role + ':', who, 10)
    para(doc, 'произвела осмотр смонтированного оборудования и составила настоящий акт о нижеследующем:', 10, before=3)
    items = [
        f'ООО «АльтЭн» выполнен монтаж оборудования {SYSG[code]} здания {bld}. Перечень, типы и количество смонтированного оборудования — по ведомости (приложение 1 к настоящему акту).',
        f'Монтаж выполнен по рабочей документации {v["rd"]}, {v["rdl"]}. Фактическое расположение оборудования, позиционные обозначения и кабельные трассы отражены в исполнительной документации {v["id"]}, {v["idl"]}, и кабельном журнале {v["kj"]}.',
        f'Монтаж выполнен в соответствии с требованиями: {NORM[code]}; эксплуатационной документацией изготовителей оборудования.',
        f'Оборудование прошло входной контроль (журнал входного контроля, том {code}, здание {bld}); паспорта изготовителей в наличии.',
        f'Скрытые работы (устройство и герметизация кабельных проходок, монтаж кабеленесущих систем, прокладка кабельных линий — в составе, предусмотренном исполнительной документацией) освидетельствованы актами № {C.aosr_list(bld, code)}.',
        'Оборудование смонтировано в полном объёме, повреждений и некомплектности не выявлено, маркировка и позиционные обозначения соответствуют исполнительной документации.',
    ]
    for k, x in enumerate(items, 1):
        p = para(doc, after=1); font(p.add_run(f'{k}. '), 10, True); font(p.add_run(x), 10, False, True)
    p = para(doc, before=4); font(p.add_run('Заключение: '), 10, True)
    font(p.add_run(f'оборудование {SYSG[code]} здания {bld} смонтировано в соответствии с исполнительной документацией и требованиями нормативных документов и готово к проведению индивидуальных испытаний (пусконаладочных работ).'), 10, False, True)
    p = para(doc, before=4); font(p.add_run('Приложения: '), 10, True)
    font(p.add_run(f'1. Ведомость смонтированного оборудования. 2. Исполнительная схема {v["id"]}, {v["idl"]}.'), 10, False, True)
    para(doc, 'Подписи:', 10, True, before=6, after=2)
    sg = doc.add_table(rows=0, cols=2)
    for role, who in [('Представитель застройщика (технического заказчика) по вопросам строительного контроля', 'Коренев Д.Ю.'),
                      ('Представитель лица, осуществляющего строительство', 'Белоногов И.Ю.'),
                      ('Представитель лица, выполнившего монтаж', 'Иванчук М.В.')]:
        rc = sg.add_row().cells
        cell(rc[0], role, 9, True); cell(rc[1], (f'    ___________________ / {who} /' if who.startswith('Белоногов') else f'____________________ / {who} /'), 10, align=WD_ALIGN_PARAGRAPH.RIGHT)
    fix_widths(sg, [Mm(110), Mm(65)])
    # ---- приложение 1: ведомость
    doc.add_page_break()
    para(doc, f'Приложение 1 к акту о смонтированном оборудовании № {no}', 10, False, True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, 'ВЕДОМОСТЬ СМОНТИРОВАННОГО ОБОРУДОВАНИЯ', 12, True, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=0)
    para(doc, f'{C.SYSN[code]}, здание {bld}. Исполнительная документация {v["id"]}, {v["idl"]}', 10, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    H = ['№', 'Наименование оборудования (по ИД)', 'Тип, марка', 'Изготовитель', 'Поз. обозн. по ИД', 'Кол-во, шт.', 'Источник в ИД']
    W = [Mm(8), Mm(62), Mm(34), Mm(20), Mm(18), Mm(13), Mm(20)]
    t = doc.add_table(rows=1, cols=len(H)); borders(t); t.autofit = False
    for k, h in enumerate(H):
        cell(t.rows[0].cells[k], h, 8, True, WD_ALIGN_PARAGRAPH.CENTER); t.rows[0].cells[k].width = W[k]
    total = 0
    for n, (pos, nm, typ, mf, qty, src) in enumerate(EQ[(bld, code)], 1):
        rr = t.add_row().cells; total += qty
        for k, x in enumerate([n, nm, typ, mf, pos, qty, src]):
            cell(rr[k], x, 8, align=WD_ALIGN_PARAGRAPH.CENTER if k in (0, 4, 5) else WD_ALIGN_PARAGRAPH.LEFT); rr[k].width = W[k]
    rr = t.add_row().cells
    cell(rr[1], 'Итого единиц оборудования', 8, True); cell(rr[5], total, 8, True, WD_ALIGN_PARAGRAPH.CENTER)
    fix_widths(t, W)
    para(doc, '', after=6)
    para(doc, 'Ведомость составил: начальник отдела систем технической безопасности ООО «АльтЭн» ____________ / Иванчук М.В. /', 10, before=6)
    os.makedirs(OUT, exist_ok=True)
    p = f'{OUT}/Акт о смонтированном оборудовании {code} {bld}.docx'; doc.save(p); return p

if __name__ == '__main__':
    for b in ('КПП', 'АБК'):
        for c in C.SYSTEMS:
            print(build(b, c))
