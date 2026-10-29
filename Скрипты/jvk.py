"""Журналы входного контроля по томам (10 шт.)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.page import PageMargins
import common as C
from equip import EQ

FN = "Times New Roman"
thin = Side(style="thin", color="000000"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
GREY = PatternFill("solid", fgColor="EDEDED")
def f(sz=10, b=False, i=False): return Font(name=FN, size=sz, bold=b, italic=i)
def al(h="left", v="top"): return Alignment(horizontal=h, vertical=v, wrap_text=True)
OUT = '/home/claude/kv/jvk'
RES_M = "Внешний осмотр, маркировка, целостность упаковки, количество — без замечаний. Документ о качестве в наличии."
RES_E = "Внешний осмотр, маркировка, комплектность — без замечаний. Паспорт изготовителя в наличии (оригинал)."
OK = "Соответствует. Допущено к применению"

def build(bld, code):
    v = C.vol(bld, code); d = C.DATES[bld]
    wb = Workbook(); ws = wb.active; ws.title = f"ЖВК {code} {bld}"
    W = [5, 11, 40, 20, 7, 8, 44, 34, 18, 16, 14]
    for k, w in enumerate(W): ws.column_dimensions[chr(65 + k)].width = w
    last = chr(64 + len(W))
    def line(rw, t, sz=10, b=False, h=None, alg="left"):
        ws.merge_cells(f"A{rw}:{last}{rw}"); c = ws[f"A{rw}"]; c.value = t; c.font = f(sz, b); c.alignment = al(alg, "center")
        if h: ws.row_dimensions[rw].height = h
    line(1, "ЖУРНАЛ", 14, True, 20, "center")
    line(2, "входного контроля строительных материалов, изделий и оборудования", 12, True, 18, "center")
    line(3, f"Том ИД: {code} — {C.SYSN[code]}. Здание: {bld}.", 11, True, 18, "center")
    line(5, f"Объект капитального строительства: {C.OBJ}", 10, h=28)
    line(6, f"Застройщик (технический заказчик): {C.ZAK}")
    line(7, f"Лицо, осуществляющее строительство: {C.STR}")
    line(8, f"Лицо, выполняющее работы: {C.ISP}")
    line(9, f"Рабочая документация {v['rd']}, {v['rdl']}; исполнительная документация {v['id']}, {v['idl']}. "
            + (f"Спецификация {v['so']}. " if v.get('so') else "") + f"Количество материалов — по АОСР № {C.aosr_list(bld, code)}; оборудование — по исполнительной документации.", 10, h=28)
    line(10, "Основание: приказ Минстроя России от 16.05.2023 № 344/пр (п. 12 состава исполнительной документации); ГОСТ 24297-2013 «Верификация закупленной продукции»; СП 48.13330.2019 «Организация строительства».", 9, h=26)
    line(11, f"Журнал начат: {d['mat']}                    Журнал окончен: {d['eq']}", 10)
    H = ["№ п/п", "Дата поступления", "Наименование продукции, марка (тип)", "Изготовитель", "Ед. изм.", "Кол-во",
         "Документ о качестве (вид, номер, дата, срок действия)", "Результаты входного контроля", "Заключение",
         "Ответственный (Ф.И.О., подпись)", "Применение"]
    hr = 13
    for k, t in enumerate(H, 1):
        c = ws.cell(hr, k, t); c.font = f(9, True); c.border = BD; c.alignment = al("center", "center"); c.fill = GREY
    ws.row_dimensions[hr].height = 44
    for k in range(1, len(H) + 1):
        c = ws.cell(hr + 1, k, k); c.font = f(8); c.border = BD; c.alignment = al("center")
    rw = hr + 2; n = 0
    def sect(t):
        nonlocal rw
        ws.merge_cells(f"A{rw}:{last}{rw}"); c = ws.cell(rw, 1, t); c.font = f(10, True)
        for k in range(1, len(H) + 1): ws.cell(rw, k).border = BD
        rw += 1
    def put(vals):
        nonlocal rw
        for k, x in enumerate(vals, 1):
            c = ws.cell(rw, k, x); c.font = f(9); c.border = BD
            c.alignment = al("center" if k in (1, 2, 5, 6) else "left")
        rw += 1
    sect("Раздел 1. Строительные материалы и изделия (по АОСР)")
    seen = {}
    for title, items in v['groups']:
        actno = title.split(' — ')[0].replace('АОСР № ', 'АОСР № ')
        for key, qty in items:
            nm, mf, un, docs = C.docs_of(key)
            if key in seen:  # один материал в нескольких актах тома — одна запись, суммарное кол-во
                i = seen[key]; ws.cell(i, 6).value = ws.cell(i, 6).value + qty
                ws.cell(i, 11).value += f"; {actno.replace('АОСР № ', '')}"
                continue
            n += 1
            dtxt = '; '.join(C.doc_text(x) for x in docs) or '—'
            seen[key] = rw
            put([n, d['mat'], nm, mf, un, qty, dtxt, RES_M, OK, "Иванчук М.В.", actno])
    sect("Раздел 2. Оборудование (по исполнительной документации)")
    for pos, nm, typ, mf, qty, src in EQ[(bld, code)]:
        n += 1
        full = nm + (f", {typ}" if typ and typ != '—' else '')
        put([n, d['eq'], full + f" (поз. {pos})" if pos != '—' else full, mf, "шт.", qty,
             "Паспорт (руководство по эксплуатации) изготовителя", RES_E, OK, "Иванчук М.В.", f"ИД {v['idl']}"])
    rw += 1
    for role, nm in [("Ответственный за входной контроль (лицо, выполняющее работы)", "начальник отдела систем технической безопасности ООО «АльтЭн» Иванчук М.В."),
                     ("Представитель лица, осуществляющего строительство", "руководитель проекта Белоногов И.Ю.")]:
        ws.merge_cells(f"A{rw}:{last}{rw}"); ws.cell(rw, 1, role).font = f(10, True); rw += 1
        ws.merge_cells(f"A{rw}:F{rw}"); ws.cell(rw, 1, nm).font = f(10, i=True)
        ws.merge_cells(f"G{rw}:H{rw}"); ws.cell(rw, 7, "____________________ (подпись)").font = f(10)
        ws.merge_cells(f"I{rw}:{last}{rw}"); ws.cell(rw, 9, "«____» ____________ 2026 г.").font = f(10)
        rw += 2
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr + 1}"; ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.print_options.horizontalCentered = True
    os.makedirs(OUT, exist_ok=True)
    p = f"{OUT}/Журнал входного контроля {code} {bld}.xlsx"; wb.save(p); return p, n

if __name__ == '__main__':
    for b in ('КПП', 'АБК'):
        for c in C.SYSTEMS:
            print(build(b, c))
