import openpyxl, random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.page import PageMargins

FN = "Times New Roman"
thin = Side(style="thin", color="000000"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
RED = "C00000"
def f(sz=10, b=False, i=False, color=None): return Font(name=FN, size=sz, bold=b, italic=i, color=color)
AL = lambda h="left", v="center", w=True: Alignment(horizontal=h, vertical=v, wrap_text=w)

# ---------- линии из кабельного журнала КПП ----------
kj = openpyxl.load_workbook('/home/claude/kzh/КЖ КПП — длины СПС 66 м.xlsx', data_only=True)['Лист 1']
rows = [r for r in kj.iter_rows(values_only=True) if r[0]]
def L(r):
    ln = r[9] if r[9] not in (None, '') else r[6]
    return dict(n=str(r[0]).strip(), a=str(r[1]).strip(), b=str(r[2]).strip(), cab=str(r[4]).strip(), sec=str(r[5]).strip(), len=float(ln))
lines = [L(r) for r in rows]
by = lambda pfx: [x for x in lines if x['n'].startswith(pfx)]

KPS = 'КПСнг(А)-FRHF 1×2×0,75'
SPS = by('ДПЛС')
SKUD_UTP = by('Д.')
P = by('P')
SKUD = SKUD_UTP + P + [dict(n='— (уточнить по КЖ)', a='по ИД, лист 3', b='', cab='КПСнг(А)-FRHF', sec='1×2×0,75', len=7.0)]
SKS = by('СКС') + by('ТФ')
SOT = by('D1')
SOTS = [dict(n='Р', a='по ИД, лист 6', b='С2000-БКИ', cab='КПСнг(А)-FRHF', sec='1×2×0,75', len=3.0)]

def volt(x):
    c = x['cab'].upper()
    if 'ПВС' in c or 'ВВГ' in c: return 1000
    if 'UTP' in c or 'CAT' in c: return 250
    return 500
def has_pe(x):
    c = x['cab'].upper(); return 'ПВС' in c or 'ВВГ' in c

HDR = ["Объект: «Центр временного содержания иностранных граждан», Калужская область, м. р-н Боровский, д. Старомихайловское, ул. 1-я Индустриальная, вл. 2",
       "Застройщик (технический заказчик): АНО «СОЮЗЭКСПЕРТИЗА» ТПП РФ. Лицо, осуществляющее строительство: ИП Белоногова О.А. Работы выполнил: ООО «АльтЭн»"]
SYSNAME = {'СПС': 'система пожарной сигнализации и СОУЭ (АПС, СОУЭ)', 'СКУД': 'система контроля и управления доступом (СКУД)',
           'СКС': 'структурированная кабельная система (СКС)', 'СОТ': 'система охранная телевизионная (СОТ)',
           'СОТС': 'система охранной и тревожной сигнализации (СОТС)'}
SHEET = {'СПС': 2, 'СКУД': 3, 'СКС': 4, 'СОТ': 5, 'СОТС': 6}

def setup(ws, widths):
    for k, w in enumerate(widths): ws.column_dimensions[chr(65 + k)].width = w
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5); ws.print_options.horizontalCentered = True

def merged(ws, r, text, last, sz=10, b=False, i=False, al="left", h=None, color=None):
    ws.merge_cells(f"A{r}:{last}{r}"); c = ws[f"A{r}"]; c.value = text; c.font = f(sz, b, i, color); c.alignment = AL(al)
    if h: ws.row_dimensions[r].height = h

def table_head(ws, r, heads, h=48):
    for k, t in enumerate(heads, 1):
        c = ws.cell(r, k, t); c.font = f(9, True); c.border = BD; c.alignment = AL("center")
    ws.row_dimensions[r].height = h
    for k in range(1, len(heads) + 1):
        c = ws.cell(r + 1, k, k); c.font = f(8); c.border = BD; c.alignment = AL("center")
    return r + 2

def put_row(ws, r, vals, center=()):
    for k, v in enumerate(vals, 1):
        c = ws.cell(r, k, v); c.font = f(9); c.border = BD; c.alignment = AL("center" if k in center else "left")

def signatures(ws, r, last, sample):
    roles = [("Измерения провёл", "________________________ (должность, Ф.И.О.)"),
             ("Представитель лица, выполнившего работы (ООО «АльтЭн»)", "начальник отдела систем технической безопасности Иванчук М.В."),
             ("Представитель лица, осуществляющего строительство", "руководитель проекта Белоногов И.Ю."),
             ("Представитель технического заказчика по вопросам строительного контроля", "главный специалист отдела строительного контроля Коренев Д.Ю.")]
    for role, nm in roles:
        merged(ws, r, role, last, 9, True); r += 1
        ws.merge_cells(f"A{r}:F{r}"); ws.cell(r, 1, nm).font = f(9, i=True)
        ws.cell(r, 8, "____________").font = f(9); r += 2
    if sample:
        merged(ws, r, "ОБРАЗЕЦ — не подписывается и в исполнительную документацию не включается", last, 11, True, al="center", color=RED)
    return r

def instrument_block(ws, r, last, kind):
    if kind == 'ins':
        items = ["Средство измерений: мегаомметр __________________ (тип), заводской № __________, свидетельство о поверке № __________ от __________, действительно до __________",
                 "Условия измерений: температура воздуха ______ °С, относительная влажность ______ %. Измерения выполнены при отключённых приборах и оборудовании (концы линий разведены)."]
    else:
        items = ["Средство измерений: кабельный тестер / анализатор __________________ (тип), заводской № __________, свидетельство о поверке (калибровке) № __________ от __________, действительно до __________",
                 "Условия измерений: температура воздуха ______ °С. Измерения выполнены на постоянной линии (permanent link) при отключённом активном оборудовании."]
    for t in items:
        merged(ws, r, t, last, 9, h=26); r += 1
    return r

def protocol_ins(wb, code, lns, sample, seed=0):
    ws = wb.create_sheet(f"{'Образец' if sample else 'Бланк'} {code} изоляция"[:31])
    W = [5, 16, 16, 18, 22, 9, 11, 13, 13, 11, 13]; setup(ws, W); last = chr(64 + len(W))
    r = 1
    if sample: merged(ws, r, "ОБРАЗЕЦ ЗАПОЛНЕНИЯ (значения — типичные для новых линий, для ориентира; в протокол вносятся только фактические показания прибора)", last, 10, True, al="center", color=RED, h=28); r += 1
    for t in HDR: merged(ws, r, t, last, 9, h=26); r += 1
    r += 1
    merged(ws, r, f"ПРОТОКОЛ № ПИ-{code}-КПП", last, 13, True, al="center", h=20); r += 1
    merged(ws, r, "измерения сопротивления изоляции кабельных линий", last, 11, True, al="center"); r += 1
    merged(ws, r, f"Здание: КПП. Система: {SYSNAME[code]}. Исполнительная документация 16/1/26-ИД-КПП, лист {SHEET[code]}; кабельный журнал 16/1/26-ИД-КПП.КЖ (раздел {code}).", last, 10, al="center", h=28); r += 1
    merged(ws, r, ("Дата измерений: «23» октября 2026 г." if sample else "Дата измерений: «____» ______________ 2026 г."), last, 10); r += 1
    r = instrument_block(ws, r, last, 'ins')
    merged(ws, r, "Нормативные требования: ПУЭ (7-е изд.), п. 1.8.37, табл. 1.8.34 — сопротивление изоляции электропроводок не менее 0,5 МОм; эксплуатационная документация изготовителей кабелей.", last, 9, h=26); r += 1
    r += 1
    H = ["№ п/п", "Обозначение линии (по КЖ)", "Начало", "Конец", "Марка кабеля", "Сечение", "Длина, м", "Напряжение мегаомметра, В",
         "R изоляции между жилами, МОм", "R изоляции жилы — PE (земля), МОм", "Норма, МОм / заключение"]
    r = table_head(ws, r, H, 54)
    rnd = random.Random(seed)
    first = r
    for k, x in enumerate(lns, 1):
        v = volt(x)
        if sample:
            lo, hi = {500: (1200, 4800), 250: (850, 3200), 1000: (1600, 6500)}[v]
            rr = round(rnd.uniform(lo, hi), -1)
            rpe = round(rnd.uniform(lo, hi), -1) if has_pe(x) else "—"
            concl = "≥ 0,5 / соответствует"
        else:
            rr = ""; rpe = "" if has_pe(x) else "—"; concl = "≥ 0,5 / "
        put_row(ws, r, [k, x['n'], x['a'], x['b'], x['cab'], x['sec'], x['len'], v, rr, rpe, concl], center=(1, 6, 7, 8, 9, 10, 11))
        r += 1
    r += 1
    merged(ws, r, "Заключение: сопротивление изоляции всех проверенных кабельных линий соответствует требованиям ПУЭ, п. 1.8.37 (не менее 0,5 МОм). Линии пригодны к подключению оборудования." if sample else
           "Заключение: ______________________________________________________________________________________________", last, 10, True, h=28); r += 2
    signatures(ws, r, last, sample)
    ws.print_title_rows = f"{first - 2}:{first - 1}"
    return ws

def protocol_lan(wb, code, lns, sample, seed=0):
    ws = wb.create_sheet(f"{'Образец' if sample else 'Бланк'} {code} витая пара"[:31])
    W = [5, 12, 10, 14, 24, 11, 11, 18, 16, 14, 13]; setup(ws, W); last = chr(64 + len(W))
    r = 1
    if sample: merged(ws, r, "ОБРАЗЕЦ ЗАПОЛНЕНИЯ (значения — типичные, для ориентира; в протокол вносятся только фактические показания тестера)", last, 10, True, al="center", color=RED, h=28); r += 1
    for t in HDR: merged(ws, r, t, last, 9, h=26); r += 1
    r += 1
    merged(ws, r, f"ПРОТОКОЛ № ПТ-{code}-КПП", last, 13, True, al="center", h=20); r += 1
    merged(ws, r, "тестирования кабельных линий на витой паре", last, 11, True, al="center"); r += 1
    merged(ws, r, f"Здание: КПП. Система: {SYSNAME[code]}. Исполнительная документация 16/1/26-ИД-КПП, лист {SHEET[code]}; кабельный журнал 16/1/26-ИД-КПП.КЖ (раздел {code}).", last, 10, al="center", h=28); r += 1
    merged(ws, r, ("Дата измерений: «19» октября 2026 г." if sample else "Дата измерений: «____» ______________ 2026 г."), last, 10); r += 1
    r = instrument_block(ws, r, last, 'lan')
    merged(ws, r, "Нормативные требования: ГОСТ Р 53245-2008 (монтаж и испытания СКС); ГОСТ Р 53246-2008. Длина постоянной линии — не более 90 м; схема разводки — по ИД (T568B).", last, 9, h=26); r += 1
    r += 1
    H = ["№ п/п", "Обозначение линии (по КЖ)", "Начало", "Конец", "Марка кабеля", "Длина по КЖ, м", "Измеренная длина, м",
         "Схема разводки (обрыв, КЗ, перепутанные / расщеплённые пары)", "Непрерывность экрана (F/UTP)", "Норма длины, м", "Результат"]
    r = table_head(ws, r, H, 66)
    rnd = random.Random(seed)
    first = r
    for k, x in enumerate(lns, 1):
        scr = x['cab'].upper().startswith('F/UTP')
        if sample:
            ml = round(x['len'] * rnd.uniform(0.95, 1.06), 1)
            wm = "Норма"; sc = "Есть" if scr else "—"; res = "Годен"
        else:
            ml = ""; wm = ""; sc = "" if scr else "—"; res = ""
        put_row(ws, r, [k, x['n'], x['a'], x['b'], x['cab'], x['len'], ml, wm, sc, "≤ 90", res], center=(1, 2, 3, 6, 7, 8, 9, 10, 11))
        r += 1
    r += 1
    merged(ws, r, "Заключение: все проверенные линии соответствуют требованиям (схема разводки — норма, длина не превышает 90 м). Линии пригодны к подключению оборудования." if sample else
           "Заключение: ______________________________________________________________________________________________", last, 10, True, h=28); r += 2
    signatures(ws, r, last, sample)
    ws.print_title_rows = f"{first - 2}:{first - 1}"
    return ws

def guide(wb):
    ws = wb.create_sheet("Как проверять цифры", 0)
    ws.column_dimensions['A'].width = 130
    T = [
     ("Как проверить протокол, который принесёт исполнитель", True),
     ("Значения в образцах — типичные для новых линий, а не замеры этого объекта. В протокол идут только показания прибора.", False),
     ("", False),
     ("Сопротивление изоляции (СПС, СКУД, СОТС)", True),
     ("• Норма — не менее 0,5 МОм. У новых кабелей длиной до 100 м обычно сотни–тысячи МОм. Значения единицы–десятки МОм на новой линии — повод разобраться (влага в коробке, повреждение оболочки, не отключён прибор).", False),
     ("• Напряжение мегаомметра: КПСнг(А)-FRHF — 500 В; витая пара в СКУД (12–24 В) — 250 В; ПВС и ВВГ (230 В) — 1000 В. Приборы и оборудование отключены, иначе испытательное напряжение их повредит.", False),
     ("• Если показание выше верхнего предела прибора, пишут «> предел», например «> 2000 МОм». Это нормально.", False),
     ("• Подозрительно: одинаковые значения во всех строках, «красивые» круглые числа подряд (1000, 1000, 1000), значения выше предела прибора без знака «>», отсутствие данных о поверке.", False),
     ("• Для двухжильных КПС и витой пары графа «жилы — PE» не заполняется (нет защитного проводника) — прочерк.", False),
     ("", False),
     ("Тестирование витой пары (СКС, СОТ)", True),
     ("• Измеренная длина обычно отличается от длины по КЖ на несколько процентов (погрешность коэффициента NVP). Совпадение до десятых во всех строках — признак переписанного КЖ, а не замера.", False),
     ("• Схема разводки — «Норма» или конкретный дефект (обрыв пары, КЗ, перепутанные пары). Дефектная линия — не «годен», её переделывают и перемеряют.", False),
     ("• Для F/UTP Cat 6 проверяется непрерывность экрана; для U/UTP — прочерк.", False),
     ("• Если у исполнителя кабельный анализатор (Fluke DSX и т. п.), к протоколу прикладывают распечатку отчёта прибора — она подтверждает цифры лучше любого бланка.", False),
     ("", False),
     ("Даты", True),
     ("• Изоляцию меряют после прокладки кабеля и до подключения оборудования: после дат АОСР на прокладку и до акта о смонтированном оборудовании.", False),
     ("• По записи заказчика протокол идёт приложением к техотчёту по итогам индивидуальных испытаний (72 ч).", False),
     ("", False),
     ("Линии в бланках взяты из кабельного журнала КПП. Линий КПС 1×2×0,75 для СКУД (7 м) в журнале нет — в бланке СКУД строка с обозначением «уточнить по КЖ». Для СОТС линия «Р» — по листу 6 ИД.", False),
    ]
    for k, (t, b) in enumerate(T, 1):
        c = ws.cell(k, 1, t); c.font = f(11 if b else 10, b); c.alignment = AL()
        if t and not b: ws.row_dimensions[k].height = 30 if len(t) > 120 else 16

def build():
    out = '/home/claude/kv/prot'
    import os; os.makedirs(out, exist_ok=True)
    wb = Workbook(); wb.remove(wb.active)
    protocol_ins(wb, 'СПС', SPS, False); protocol_ins(wb, 'СКУД', SKUD, False)
    protocol_lan(wb, 'СКС', SKS, False); protocol_lan(wb, 'СОТ', SOT, False); protocol_ins(wb, 'СОТС', SOTS, False)
    b1 = f'{out}/Протоколы замеров КПП — бланки.xlsx'; wb.save(b1)
    wb = Workbook(); wb.remove(wb.active)
    guide(wb)
    protocol_ins(wb, 'СПС', SPS, True, 11); protocol_ins(wb, 'СКУД', SKUD, True, 22); protocol_lan(wb, 'СОТ', SOT, True, 33)
    b2 = f'{out}/Протоколы замеров КПП — образцы заполнения.xlsx'; wb.save(b2)
    print(b1, b2, len(SPS), len(SKUD), len(SKS), len(SOT))

build()
