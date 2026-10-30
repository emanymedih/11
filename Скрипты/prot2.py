import re
"""Протоколы испытаний по томам: бланк + образец (изоляция / витая пара)."""
import os, random, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.page import PageMargins
import common as C

FN = "Times New Roman"; RED = "C00000"
thin = Side(style="thin", color="000000"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
GREY = PatternFill("solid", fgColor="EDEDED")
def f(sz=10, b=False, i=False, color=None): return Font(name=FN, size=sz, bold=b, italic=i, color=color)
def al(h="left", v="center"): return Alignment(horizontal=h, vertical=v, wrap_text=True)
import paths
OUT = str(paths.TMP)

# ---------------- линии из кабельных журналов ----------------
def L(n, a, b, cab, sec, ln): return dict(n=str(n).strip(), a=str(a).strip(), b=str(b).strip(), cab=str(cab).strip(), sec=str(sec).strip(), len=float(ln))
def kpp_lines():
    ws = openpyxl.load_workbook(paths.DATA / 'КЖ_КПП_итог.xlsx', data_only=True)['Лист 1']
    sec_map = {'АПС СОУЭ': 'СПС', 'СКС': 'СКС', 'СКУД': 'СКУД', 'СОТ': 'СОТ', 'СОТС': 'СОТС'}
    out = {k: [] for k in sec_map.values()}; sec = None
    ren = {'EMH19.1': 'EMG19.1', 'EMH20.1': 'EMG20.1', 'EMH20.2': 'EMG20.2'}
    for r in ws.iter_rows(values_only=True):
        v = list(r[:10])
        if v[0] and all(c in (None, '') for c in v[1:]): sec = sec_map.get(str(v[0]).strip()); continue
        if not v[0] or not sec: continue
        ln = v[9] if isinstance(v[9], (int, float)) else v[6]
        cab = str(v[4]).strip()
        if 'UTP-2P' in cab: cab = 'ParLan U/UTP Cat 5e 4×2×0,52 PVC'
        if cab.startswith('F/UTP Cat 6'): cab = 'Cabeus F/UTP Cat 6 LSZH'
        sec_ = str(v[5]).strip()
        v[2] = re.sub(r'^Xs', 'XS', str(v[2]).strip()) if v[2] else v[2]
        out[sec].append(L(v[0], ren.get(str(v[1]).strip(), v[1]), ren.get(str(v[2]).strip(), v[2]), cab, sec_, ln))
    return out
def abk_lines():
    ws = openpyxl.load_workbook(paths.DATA / 'КЖ_АБК.xlsx', data_only=True)['Лист1']
    out = {}; sys = None
    for r in ws.iter_rows(values_only=True):
        v = list(r[:10])
        if v[0] and all(c is None for c in v[1:]):
            sys = {'АПС,СОУЭ': 'СПС'}.get(str(v[0]).strip(), str(v[0]).strip()); out[sys] = []; continue
        if not v[0] and v[1] and out[sys]: v[0] = out[sys][-1]['n']  # ветка без обозначения в КЖ — по предыдущей линии (ДПЛС2.1)
        if not v[0]: continue
        ln = v[9] if isinstance(v[9], (int, float)) else v[6]
        out[sys].append(L(v[0], v[1], v[2], v[4], v[5], ln))
    return out
LINES = {'КПП': kpp_lines(), 'АБК': abk_lines()}

def kind(x):
    c = (x['cab'] + ' ' + x['sec']).upper(); n = x['n'].upper()
    if 'ПВС' in c or 'ВВГ' in c: return 'LV'          # 230 В
    if n.startswith('V'): return 'V100'               # речевое оповещение 100 В
    return 'SELV'                                      # 12–24 В, шлейфы, интерфейсы, витая пара СКУД
TEST = {'SELV': (250, '0,5'), 'V100': (500, '1,0'), 'LV': (500, '1,0')}
RANGE = {'SELV': (600, 3800), 'V100': (900, 5200), 'LV': (1200, 6800)}
LAN = set()
NOBANNER = set()  # плашки «ОБРАЗЕЦ» во всех протоколах (решение пользователя 30.10.2026)

HDR_LAB = ["ООО «АльтЭн» (лицо, выполнившее работы), ОГРН 1024000944098, ИНН 4025062969",
           "249035, Калужская обл., г. Обнинск, пр-т Маркса, д. 14, пом. 18"]
MONTHS = ["", "января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря"]

def setup(ws, W):
    for k, w in enumerate(W): ws.column_dimensions[chr(65 + k)].width = w
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5); ws.print_options.horizontalCentered = True

def protocol(wb, bld, code, sample, seed):
    lan = code in LAN; lines = LINES[bld][code]; v = C.vol(bld, code); dt = C.DATES[bld]['prot']
    ws = wb.create_sheet(('Образец' if sample else 'Бланк'))
    W = ([5, 12, 11, 15, 22, 9, 9, 15, 10, 10, 10, 11] if lan else [5, 15, 14, 18, 22, 10, 8, 11, 13, 13, 10, 13])
    setup(ws, W); last = chr(64 + len(W)); r = 1
    def line(t, sz=10, b=False, alg="left", h=None, color=None):
        nonlocal r
        ws.merge_cells(f"A{r}:{last}{r}"); c = ws[f"A{r}"]; c.value = t; c.font = f(sz, b, False, color); c.alignment = al(alg)
        if h: ws.row_dimensions[r].height = h
        r += 1
    if sample and (bld, code) not in NOBANNER:
        line("ОБРАЗЕЦ ЗАПОЛНЕНИЯ — для сверки с протоколом исполнителя измерений. Значения типичны для новых линий данного типа и длины; "
             "в действительном протоколе — только фактические показания прибора. Не подписывается и в ИД не включается.", 10, True, "center", 30, RED)
    for t in HDR_LAB: line(t, 10, t.startswith('ООО'), "center")
    r += 1
    num = f"№ ПИ-{code}-{bld}"
    title = "измерения сопротивления изоляции кабельных линий" if not lan else "испытаний кабельных линий структурированной кабельной системы (постоянная линия)"
    line(f"ПРОТОКОЛ {num}", 13, True, "center", 20)
    line(title, 11, True, "center")
    d = dt if sample else "__.__.____"
    line(f"Дата проведения испытаний: {d}", 10)
    line(f"Объект: {C.OBJ}", 9, h=26)
    line(f"Здание: {bld}. Система: {C.SYSN[code]}. Заказчик испытаний: {C.ISP}.", 9)
    line(f"Исполнительная документация: {v['id']}, {v['idl']}; кабельный журнал {v['kj']}. Рабочая документация {v['rd']}.", 9, h=26)
    line("Цель испытаний: " + ("приёмо-сдаточные испытания — проверка сопротивления изоляции смонтированных кабельных линий." if not lan
         else "приёмо-сдаточные испытания — сертификационное тестирование постоянных линий (permanent link) медной подсистемы."), 9)
    cond = "температура воздуха +20 °С, относительная влажность 55 %, атмосферное давление 745 мм рт. ст." if sample else "температура воздуха ____ °С, относительная влажность ____ %, атмосферное давление ____ мм рт. ст."
    line(f"Условия испытаний: {cond}", 9)
    if not lan:
        line("Нормативные документы: ГОСТ Р 50571.16-2019 (МЭК 60364-6:2016), п. 6.4.3.3; ПУЭ (7-е изд.), п. 1.8.37; СП 76.13330.2016; эксплуатационная документация изготовителей кабелей.", 9, h=26)
        line("Нормы: цепи SELV/PELV (до 50 В: шлейфы, интерфейсы, питание 12–24 В) — испытательное напряжение 250 В, сопротивление изоляции не менее 0,5 МОм; "
             "цепи до 500 В (линии оповещения 100 В, сеть 230 В) — 500 В, не менее 1,0 МОм. Измеряется сопротивление изоляции между жилами и между жилами и землёй (ГОСТ Р 50571.16-2019, п. 6.4.3.4: для SELV — отделение от земли); для кабелей без PE-жилы — относительно экрана (F/UTP) или заземлённой конструкции (лоток, шина PE шкафа). Измерения при отключённых приборах (концы линий разведены).", 9, h=52)
        inst = ("Измеритель параметров электробезопасности Sonel MPI-525 (Госреестр СИ РФ № 44170-10), зав. № A92551; "
                "поверка: ФГИС «Аршин», свидетельство № С-ВВО/28-05-2026/529221088 от 28.05.2026, действительна до 27.05.2027. "
                "Режим RISO (ГОСТ IEC 61557-2): 250 В — диапазон 0,25…999 МОм; 500 В — 0,5 МОм…2,00 ГОм; измерительный ток < 2 мА. "
                "Измерения выполнил: инженер Андреев М.В., III группа по электробезопасности до 1000 В.")
    else:
        line("Нормативные документы: ГОСТ Р 53245-2008; ГОСТ Р 53246-2008; ANSI/TIA-568.2-D (пределы испытаний постоянной линии); СП 76.13330.2016.", 9, h=26)
        line("Пределы испытаний: TIA Cat 6 Permanent Link — для F/UTP кат. 6; TIA Cat 5e Permanent Link — для U/UTP кат. 5e. Длина постоянной линии не более 90 м. "
             "Результат «PASS» — все параметры в пределах нормы с положительным запасом.", 9, h=26)
        inst = ("Кабельный анализатор Fluke Networks DSX-5000 (пример типа), зав. № [______], адаптеры постоянной линии DSX-PLA004S, "
                "калибровка/поверка № [______] до [__.__.____]; версия ПО/база пределов — по отчёту прибора" if sample
                else "Кабельный анализатор ______________________, зав. № ________, калибровка/поверка № __________ до ________")
    line(f"Средства измерений: {inst}", 9, h=40)
    r += 1
    if not lan:
        H = ["№ п/п", "Обозначение линии (по КЖ)", "Начало", "Конец", "Марка кабеля", "Сечение", "Длина, м", "Испыт. напряжение, В",
             "R изоляции между жилами, МОм", "R изоляции жил относительно земли (PE / экран / заземлённый лоток), МОм", "Норма, не менее, МОм", "Вывод"]
    else:
        H = ["№ п/п", "Линия (по КЖ)", "Начало", "Конец", "Кабель", "Длина по КЖ, м", "Длина измер., м", "Предел испытаний",
             "Запас NEXT, дБ", "Запас RL, дБ", "Запас IL, дБ", "Результат"]
    for k, t in enumerate(H, 1):
        c = ws.cell(r, k, t); c.font = f(8, True); c.border = BD; c.alignment = al("center"); c.fill = GREY
    ws.row_dimensions[r].height = 46; hr = r; r += 1
    rnd = random.Random(seed)
    for n, x in enumerate(lines, 1):
        if not lan:
            kd = kind(x); uv, norm = TEST[kd]; pe = True
            if sample:
                top = 999 if uv == 250 else 2000
                def meas():
                    if rnd.random() < 0.8: return f"> {top}" if uv == 250 else "> 2,00 ГОм"
                    v_ = int(round(rnd.uniform(0.35, 0.95) * top, -1)); return v_
                rr = meas(); rpe = meas() if pe else "—"; res = "соответствует"
            else:
                rr = ""; rpe = "" if pe else "—"; res = ""
            vals = [n, x['n'], x['a'], x['b'], x['cab'], x['sec'], x['len'], uv, rr, rpe, norm, res]
        else:
            cat6 = 'CAT.6' in x['cab'].upper() or 'CAT 6' in x['cab'].upper()
            lim = "TIA Cat 6 PL" if cat6 else "TIA Cat 5e PL"
            if sample:
                ml = round(x['len'] * rnd.uniform(0.96, 1.05), 1)
                nx = round(rnd.uniform(4.0, 10.5) if cat6 else rnd.uniform(6.5, 13.5), 1)
                rl = round(rnd.uniform(3.0, 8.5) if cat6 else rnd.uniform(4.0, 10.0), 1)
                il = round(((30.7 - (0.33 * ml + 0.5)) if cat6 else (21.0 - (0.22 * ml + 0.4))) + rnd.uniform(-0.8, 0.8), 1)
                res = "PASS"
            else:
                ml = nx = rl = il = res = ""
            vals = [n, x['n'], x['a'], x['b'], x['cab'], x['len'], ml, lim, nx, rl, il, res]
        for k, val in enumerate(vals, 1):
            c = ws.cell(r, k, val); c.font = f(8); c.border = BD; c.alignment = al("left" if k in (2, 3, 4, 5) else "center")
        r += 1
    r += 1
    if sample:
        concl = ("Сопротивление изоляции всех испытанных кабельных линий соответствует требованиям ГОСТ Р 50571.16-2019 и ПУЭ п. 1.8.37. Линии пригодны к эксплуатации."
                 if not lan else "Все испытанные постоянные линии соответствуют пределам испытаний (PASS). Подсистема пригодна к эксплуатации. Отчёт кабельного анализатора (PDF/LinkWare) — приложение к протоколу.")
    else:
        concl = "______________________________________________________________________________________________"
    line(f"Заключение: {concl}", 10, True, h=28)
    line("Протокол распространяется только на испытанные линии. Запись «> 999» / «> 2,00 ГОм» — показание выше верхнего предела диапазона прибора.", 8)
    r += 1
    for role, nm in [("Измерения выполнил", "инженер Андреев М.В. (III группа по ЭБ до 1000 В)"),
                     ("Протокол проверил", "начальник отдела систем технической безопасности Иванчук М.В."),
                                          ("Присутствовал представитель лица, осуществляющего строительство", "руководитель проекта Белоногов И.Ю.")]:
        ws.merge_cells(f"A{r}:F{r}"); ws.cell(r, 1, role).font = f(9, True)
        ws.merge_cells(f"G{r}:J{r}"); ws.cell(r, 7, nm).font = f(9, i=True)
        ws.merge_cells(f"K{r}:{last}{r}"); ws.cell(r, 11, "__________").font = f(9)
        r += 1
    if sample and (bld, code) not in NOBANNER:
        r += 1; line("ОБРАЗЕЦ", 14, True, "center", color=RED)
    ws.print_title_rows = f"{hr}:{hr}"
    return len(lines)

def build(bld, code):
    wb = Workbook(); wb.remove(wb.active)
    n = protocol(wb, bld, code, True, sum(map(ord, bld + code)))
    protocol(wb, bld, code, False, 0)
    os.makedirs(OUT, exist_ok=True)
    kind_ = 'тестирования линий' if code in LAN else 'сопротивления изоляции'
    p = f"{OUT}/Протокол {kind_} {code} {bld} (образец и бланк).xlsx"; wb.save(p); return p, n

if __name__ == '__main__':
    for b in ('КПП', 'АБК'):
        for c in C.SYSTEMS:
            print(build(b, c))
