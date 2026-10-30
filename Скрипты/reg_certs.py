"""Сводный реестр документов о качестве, прикладываемых к ИД здания (КПП / АБК, по всем томам)."""
import openpyxl, collections, paths
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
import aso, common as C
from equip import EQ

SYS = ['СПС', 'СКУД', 'СКС', 'СОТ', 'СОТС']


def build(bld):
    docs = collections.OrderedDict()
    for s in SYS:
        ws = openpyxl.load_workbook(paths.OUT / bld / f'Том {s} {bld}' / f'Реестр документов о качестве {s} {bld}.xlsx').active
        for r in ws.iter_rows(values_only=True):
            if not (isinstance(r[0], int) and r[1] != 2): continue
            pos, mat, mf, un, q, kind, num, dt, srok, who, lst, act = r[:12]
            k = (kind, num, dt, who)
            d = docs.setdefault(k, {'kind': kind, 'num': num, 'dt': dt, 'srok': srok, 'who': who, 'lst': lst, 'mat': [], 'pos': {}})
            m = f'{mat} ({mf})'
            if m not in d['mat']: d['mat'].append(m)
            d['pos'].setdefault(s, []).append(pos)

    NOTE = {
     'ЕАЭС RU C-RU.0010.В.00450/24': 'Копия не получена — запросить у поставщика.',
     'РОСС CN.НВ61.Н29435': f"Добровольный. Действовал на дату входного контроля ({C.DATES[bld]['mat']}). Принят заказчиком.",
    }

    doc = Document()
    s = doc.sections[0]; s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Mm(297), Mm(210)
    s.left_margin = s.right_margin = Mm(10); s.top_margin = s.bottom_margin = Mm(10)
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(10)
    aso.para(doc, 'СВОДНЫЙ РЕЕСТР', 13, True, align=A.CENTER, after=0)
    aso.para(doc, f'документов о качестве (сертификатов, деклараций, паспортов, писем) на материалы и изделия, прикладываемых к исполнительной документации здания {bld}', 11, True, align=A.CENTER, after=4)
    aso.para(doc, f'Объект: {C.OBJ}', 9, after=0)
    aso.para(doc, f'Застройщик (технический заказчик): {C.ZAK}', 9, after=0)
    aso.para(doc, f'Лицо, осуществляющее строительство: {C.STR}', 9, after=0)
    aso.para(doc, f'Лицо, выполнившее работы: {C.ISP}', 9, after=4)
    aso.para(doc, f'Назначение: для сверки комплектности документов о качестве, приложенных к исполнительной документации здания {bld}. Копии заверяются лицом, выполнившим работы.', 9, False, True, after=4)

    aso.para(doc, 'Раздел 1. Документы о качестве на материалы и изделия (сертификаты, решения, письма, заключения, паспорта)', 11, True, after=2)
    H = ['№', 'Вид документа', 'Номер', 'Дата выдачи', 'Срок действия', 'Кем выдан', 'Продукция', 'Кол-во листов', 'Примечание']
    W = [8, 42, 34, 18, 24, 40, 64, 13, 34]
    t = doc.add_table(rows=1, cols=len(H)); aso.borders(t)
    for k, h in enumerate(H): aso.cell(t.rows[0].cells[k], h, 8, True, A.CENTER)
    tl = 0
    for n, d in enumerate(docs.values(), 1):
        c = t.add_row().cells
        tl += d['lst'] or 0
        vals = [n, d['kind'], d['num'], d['dt'], d['srok'], d['who'], '; '.join(d['mat']), d['lst'], NOTE.get(d['num'], '')]
        for k, v in enumerate(vals):
            aso.cell(c[k], '' if v is None else str(v), 8, align=A.CENTER if k in (0, 3, 7) else A.LEFT)
    c = t.add_row().cells
    aso.cell(c[1], f'Итого: документов — {len(docs)}', 8, True)
    aso.cell(c[7], str(tl), 8, True, A.CENTER)
    aso.fix_widths(t, [Mm(w) for w in W])
    aso.para(doc, '', after=2)
    aso.para(doc, 'Раздел 2. Паспорта (руководства по эксплуатации) изготовителей на оборудование и декларации', 11, True, before=6, after=2)
    H2 = ['№', 'Оборудование, тип/марка', 'Изготовитель', 'Кол-во, шт.', 'Документ']
    W2 = [8, 145, 40, 18, 66]
    t2 = doc.add_table(rows=1, cols=len(H2)); aso.borders(t2)
    for k, h in enumerate(H2): aso.cell(t2.rows[0].cells[k], h, 8, True, A.CENTER)
    eq = collections.OrderedDict(); by_typ = {}
    for x in SYS:
        for pos, nm, typ, mf, qty, src in EQ[(bld, x)]:
            key = (nm.split(' (в ')[0] + (f', {typ}' if typ != '—' else ''), mf)
            t0 = typ.split(' (')[0].replace('ШПС-24', 'ШПС-24 исп.12').replace('исп.12 исп.12', 'исп.12')
            if typ != '—': key = by_typ.setdefault(t0, key)  # одно изделие в разных томах — одной строкой
            eq[key] = eq.get(key, 0) + qty
    for n2, ((nm, mf), qty) in enumerate(eq.items(), 1):
        dd = 'Паспорт изготовителя'
        if 'Delta' in mf: dd = 'Паспорт изготовителя; декларация о соответствии РОСС RU Д-VN.РА01.В.24646/23'
        c = t2.add_row().cells
        for k, v in enumerate([n2, nm, mf, qty, dd]):
            aso.cell(c[k], str(v), 8, align=A.CENTER if k in (0, 3) else A.LEFT)
    aso.fix_widths(t2, [Mm(w) for w in W2])
    aso.para(doc, '', after=8)
    sg = doc.add_table(rows=0, cols=2)
    for role, who in [('Составил: начальник отдела систем технической безопасности ООО «АльтЭн»', 'Иванчук М.В.'),
                      ('Проверил: представитель лица, осуществляющего строительство, руководитель проекта', 'Белоногов И.Ю.')]:
        rc = sg.add_row().cells
        aso.cell(rc[0], role, 9, True); aso.cell(rc[1], f'____________________ / {who} /     «____» ____________ 2026 г.', 9, align=A.RIGHT)
    aso.fix_widths(sg, [Mm(150), Mm(127)])
    out = str(paths.OUT / bld / f'Сводный реестр документов о качестве {bld}.docx')
    doc.save(out); print(out, len(docs), tl, len(eq))
    return docs, eq
